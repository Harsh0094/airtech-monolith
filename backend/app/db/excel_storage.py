
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DATA_DIR = PROJECT_ROOT / "data"
EXCEL_FILE = DATA_DIR / "airtech.xlsx"


# Each worksheet has its own columns.
SHEETS = {
    "Customers": [
        "customer_id",
        "name",
        "phone",
        "email",
        "address",
    ],
    "Quotations": [
        "quotation_id",
        "quotation_number",
        "customer_name",
        "customer_company",
        "contact_person",
        "customer_email",
        "customer_phone",
        "customer_address",
        "site_name",
        "site_address",
        "category",
        "quotation_date",
        "validity_days",
        "subtotal",
        "tax_type",
        "tax_percent",
        "tax_amount",
        "discount",
        "total",
        "status",
        "frequency_details",
        "contract_months",
        "maintenance_coverage",
        "notes",
        "terms",
        "terms_template_id",
        "terms_template_name",
        "terms_template_content_html",
        "terms_template_include_in_pdf",
    ],
    "QuotationItems": [
        "quotation_id",
        "description",
        "product_code",
        "quantity",
        "unit",
        "rate",
        "rate_unit",
        "duration",
        "line_total",
    ],
    "Rentals": [
        "rental_id",
        "customer_id",
        "equipment",
        "quantity",
        "monthly_rate",
        "months",
        "status",
    ],
    "Airtime": [
        "airtime_id",
        "customer_id",
        "service",
        "quantity",
        "amount",
        "status",
    ],
    "TermsTemplates": [
        "template_id",
        "name",
        "company",
        "status",
        "content_html",
        "include_in_pdf",
        "created_at",
        "updated_at",
    ],
}


def ensure_workbook() -> Path:
    """Create the workbook if needed and safely add missing worksheets/headers."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    if not EXCEL_FILE.exists():
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Customers"

        try:
            for sheet_name, headers in SHEETS.items():
                if sheet_name == "Customers":
                    current_sheet = worksheet
                else:
                    current_sheet = workbook.create_sheet(sheet_name)

                current_sheet.append(headers)

            workbook.save(EXCEL_FILE)
        finally:
            workbook.close()

        return EXCEL_FILE

    workbook = load_workbook(EXCEL_FILE)

    try:
        for sheet_name, headers in SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                worksheet = workbook.create_sheet(sheet_name)
                worksheet.append(headers)
                continue

            worksheet = workbook[sheet_name]
            existing_headers = [
                cell.value for cell in worksheet[1]
            ]

            if not any(header is not None for header in existing_headers):
                worksheet.append(headers)
                continue

            # Preserve existing columns and append missing ones.
            for header in headers:
                if header not in existing_headers:
                    existing_headers.append(header)
                    worksheet.cell(
                        row=1,
                        column=len(existing_headers),
                        value=header,
                    )

        workbook.save(EXCEL_FILE)
    finally:
        workbook.close()

    return EXCEL_FILE


def read_records(sheet_name: str) -> list[dict[str, Any]]:
    """Read worksheet rows as dictionaries."""
    if sheet_name not in SHEETS:
        raise ValueError(f"Unknown worksheet: {sheet_name}")

    ensure_workbook()
    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, ())

        return [
            dict(zip(headers, row))
            for row in rows
            if any(value is not None for value in row)
        ]
    finally:
        workbook.close()


def append_record(
    sheet_name: str,
    record: dict[str, Any],
) -> dict[str, Any]:
    """Append a record using the worksheet's actual headers."""
    if sheet_name not in SHEETS:
        raise ValueError(f"Unknown worksheet: {sheet_name}")

    expected_headers = SHEETS[sheet_name]
    unknown_fields = set(record) - set(expected_headers)

    if unknown_fields:
        raise ValueError(
            f"Unknown fields for {sheet_name}: "
            f"{', '.join(sorted(unknown_fields))}"
        )

    ensure_workbook()
    workbook = load_workbook(EXCEL_FILE)

    try:
        worksheet = workbook[sheet_name]
        actual_headers = [
            cell.value for cell in worksheet[1]
            if cell.value is not None
        ]

        for header in expected_headers:
            if header not in actual_headers:
                actual_headers.append(header)
                worksheet.cell(
                    row=1,
                    column=len(actual_headers),
                    value=header,
                )

        saved_record = {
            header: record.get(header)
            for header in expected_headers
        }

        worksheet.append([
            saved_record.get(header)
            for header in actual_headers
        ])

        workbook.save(EXCEL_FILE)
    finally:
        workbook.close()

    return saved_record


def delete_quotation_records(quotation_id: str) -> bool:
    """Delete a quotation and its associated line items from Excel."""
    ensure_workbook()
    workbook = load_workbook(EXCEL_FILE)

    try:
        quotations_sheet = workbook["Quotations"]
        items_sheet = workbook["QuotationItems"]

        quotation_headers = [
            cell.value for cell in quotations_sheet[1]
        ]
        item_headers = [
            cell.value for cell in items_sheet[1]
        ]

        quotation_id_col = quotation_headers.index("quotation_id") + 1
        item_id_col = item_headers.index("quotation_id") + 1

        # Find the quotation row by its exact ID.
        quotation_row = next(
            (
                row
                for row in range(2, quotations_sheet.max_row + 1)
                if str(
                    quotations_sheet.cell(
                        row=row,
                        column=quotation_id_col,
                    ).value or ""
                ) == quotation_id
            ),
            None,
        )

        if quotation_row is None:
            return False

        # Delete associated line items from bottom to top.
        for row in range(items_sheet.max_row, 1, -1):
            stored_id = str(
                items_sheet.cell(
                    row=row,
                    column=item_id_col,
                ).value or ""
            )

            if stored_id == quotation_id:
                items_sheet.delete_rows(row, 1)

        quotations_sheet.delete_rows(quotation_row, 1)
        workbook.save(EXCEL_FILE)
        return True
    finally:
        workbook.close()
