
from datetime import datetime, timezone
from uuid import uuid4

from fastapi import APIRouter, HTTPException

from app.db.excel_storage import append_record, read_records

router = APIRouter()


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


@router.get("/")
def list_terms_templates():
    """Return all saved terms templates."""
    templates = read_records("TermsTemplates")
    return {"templates": templates}


@router.post("/")
def create_terms_template(payload: dict):
    """Create a reusable terms template."""
    name = str(payload.get("name", "")).strip()
    company = str(payload.get("company", "Airtech")).strip()
    content_html = str(payload.get("content_html", "")).strip()
    status = str(payload.get("status", "Draft")).strip()
    include_in_pdf = payload.get("include_in_pdf", True)

    if not name:
        raise HTTPException(
            status_code=422,
            detail="Template name is required.",
        )

    if not content_html:
        raise HTTPException(
            status_code=422,
            detail="Template content is required.",
        )

    if status not in {"Draft", "Active", "Archived"}:
        raise HTTPException(
            status_code=422,
            detail="Status must be Draft, Active, or Archived.",
        )

    if not isinstance(include_in_pdf, bool):
        raise HTTPException(
            status_code=422,
            detail="include_in_pdf must be true or false.",
        )

    templates = read_records("TermsTemplates")

    if any(
        str(template.get("name", "")).strip().casefold()
        == name.casefold()
        for template in templates
    ):
        raise HTTPException(
            status_code=409,
            detail="A template with this name already exists.",
        )

    timestamp = now_iso()

    record = {
        "template_id": uuid4().hex,
        "name": name,
        "company": company or "Airtech",
        "status": status,
        "content_html": content_html,
        "include_in_pdf": include_in_pdf,
        "created_at": timestamp,
        "updated_at": timestamp,
    }

    saved = append_record("TermsTemplates", record)
    return {"template": saved}


@router.put("/{template_id}")
def update_terms_template(template_id: str, payload: dict):
    """Update an existing template without changing its ID or creation time."""
    templates = read_records("TermsTemplates")
    existing = next(
        (
            template for template in templates
            if str(template.get("template_id")) == template_id
        ),
        None,
    )

    if existing is None:
        raise HTTPException(
            status_code=404,
            detail="Terms template not found.",
        )

    updated = dict(existing)

    for field in ("name", "company", "content_html", "status"):
        if field in payload:
            updated[field] = str(payload[field]).strip()

    if "include_in_pdf" in payload:
        if not isinstance(payload["include_in_pdf"], bool):
            raise HTTPException(
                status_code=422,
                detail="include_in_pdf must be true or false.",
            )
        updated["include_in_pdf"] = payload["include_in_pdf"]

    if not updated.get("name"):
        raise HTTPException(
            status_code=422,
            detail="Template name is required.",
        )

    if not updated.get("content_html"):
        raise HTTPException(
            status_code=422,
            detail="Template content is required.",
        )

    if updated.get("status") not in {"Draft", "Active", "Archived"}:
        raise HTTPException(
            status_code=422,
            detail="Status must be Draft, Active, or Archived.",
        )

    duplicate = any(
        str(template.get("template_id")) != template_id
        and str(template.get("name", "")).strip().casefold()
        == updated["name"].casefold()
        for template in templates
    )

    if duplicate:
        raise HTTPException(
            status_code=409,
            detail="A different template already uses this name.",
        )

    updated["updated_at"] = now_iso()

    # Rewrite the worksheet row while preserving other templates.
    from openpyxl import load_workbook
    from app.db.excel_storage import EXCEL_FILE, ensure_workbook, SHEETS

    ensure_workbook()
    workbook = load_workbook(EXCEL_FILE)

    try:
        worksheet = workbook["TermsTemplates"]
        headers = [cell.value for cell in worksheet[1]]
        row_number = None

        for row in range(2, worksheet.max_row + 1):
            row_id = worksheet.cell(
                row=row,
                column=headers.index("template_id") + 1,
            ).value

            if str(row_id) == template_id:
                row_number = row
                break

        if row_number is None:
            raise HTTPException(
                status_code=404,
                detail="Terms template not found.",
            )

        for column, header in enumerate(headers, start=1):
            if header in SHEETS["TermsTemplates"]:
                worksheet.cell(
                    row=row_number,
                    column=column,
                    value=updated.get(header),
                )

        workbook.save(EXCEL_FILE)
    finally:
        workbook.close()

    return {"template": updated}



@router.delete("/{template_id}")
def delete_terms_template(template_id: str):
    """Delete one reusable terms template without modifying existing quotations."""
    from openpyxl import load_workbook
    from app.db.excel_storage import EXCEL_FILE, ensure_workbook

    ensure_workbook()
    workbook = load_workbook(EXCEL_FILE)

    try:
        worksheet = workbook["TermsTemplates"]
        headers = [cell.value for cell in worksheet[1]]

        if "template_id" not in headers:
            raise HTTPException(
                status_code=500,
                detail="The template_id column is missing.",
            )

        id_column = headers.index("template_id") + 1
        row_number = None

        for row in range(2, worksheet.max_row + 1):
            row_id = worksheet.cell(
                row=row,
                column=id_column,
            ).value

            if str(row_id) == template_id:
                row_number = row
                break

        if row_number is None:
            raise HTTPException(
                status_code=404,
                detail="Terms template not found.",
            )

        worksheet.delete_rows(row_number, 1)
        workbook.save(EXCEL_FILE)

    finally:
        workbook.close()

    return {
        "message": "Terms template deleted successfully.",
        "template_id": template_id,
    }
